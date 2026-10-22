"""Export the R-generated SVG and populate the approved Word table layout.
Only aggregate counts are exported. Existing individual-level inputs stay local.
"""
from pathlib import Path
from io import BytesIO
import sys, copy, hashlib, json, re
import pandas as pd
from PIL import Image
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

ROOT=Path(sys.argv[1]) if len(sys.argv)>1 else Path('Z:/Research/MSI-1')
BASE=ROOT/'analysis_updates/FigureS5_organ_support_20261019'
OUT=BASE/'outputs'; QA=BASE/'qa'
sys.path.insert(0,str(ROOT/'manuscript_20261019/submission_build/vendor'))
import resvg_py

FIGNAME='Figure S5. Trajectory classes, early organ support and MAKE30'
TABNAME='Table S6. Early organ support and MAKE30 by trajectory class'
svg=(OUT/(FIGNAME+'.svg')).read_text(encoding='utf-8')
for w,suffix in [(2700,'.png'),(6600,'.tif')]:
    im=Image.open(BytesIO(resvg_py.svg_to_bytes(svg_string=svg,width=w,background='white',font_files=['C:/Windows/Fonts/arial.ttf','C:/Windows/Fonts/arialbd.ttf']))).convert('RGB')
    if suffix=='.tif':im.save(OUT/(FIGNAME+suffix),compression='tiff_lzw',dpi=(600,600))
    else:im.save(OUT/(FIGNAME+suffix),dpi=(245.45,245.45))

summary=pd.read_csv(OUT/'Support_and_MAKE30_summary.csv')
counts=pd.read_csv(OUT/'Sankey_path_counts.csv')
audit=pd.read_csv(OUT/'Input_and_cohort_audit.csv')
# Windows R in the C locale may escape non-ASCII path characters in CSV output.
audit['source']=audit['source'].map(lambda s: re.sub(r'<U\+([0-9A-Fa-f]+)>',lambda m: chr(int(m.group(1),16)),s))
audit.to_csv(OUT/'Input_and_cohort_audit.csv',index=False,encoding='utf-8-sig')
labels={'MIMIC':'MIMIC-IV','eICU':'eICU-CRD','SICdb':'SICdb'}
classes=['C1','C2','C3']; dbs=list(labels)
supports=['Neither support','Vasopressors only','Mechanical ventilation only','Both supports']

# Independent reconciliation against source records and the R aggregate tables.
for db in dbs:
    a=audit[audit.cohort==db].iloc[0]
    f=Path(a.source);assert hashlib.sha256(f.read_bytes()).hexdigest()==a.source_sha256
    d=pd.read_csv(f);d=d[d.make30_analysis.isin([0,1])].copy()
    assert d.cohort_unit_id.is_unique and len(d)==a.analysis_n
    d['support']=d.vasopressor_0_24h_analysis+2*d.mechanical_ventilation_0_24h_analysis
    for _,r in counts[counts.cohort==db].iterrows():
        mask=(d.assigned_class==int(r['class'][1]))&(d.support==r.support)&(d.make30_analysis==r.event)
        assert int(mask.sum())==r['n']
    for s in range(4):
        r=summary[(summary.cohort==db)&(summary.support==s)&summary['class'].isin(classes)].set_index('class')
        assert r.loc['C2','event_percent']>r.loc['C3','event_percent']>r.loc['C1','event_percent']

template=ROOT/'JoGH_Submission_20261019/Supplementary materials/Table S2. MAKE30 and component outcomes by trajectory class.docx'
template_hash=hashlib.sha256(template.read_bytes()).hexdigest()
doc=Document(template)
original_table=copy.deepcopy(doc.tables[0]._tbl)
# Preserve the authoritative section, styles, header/footer parts; replace body slots.
for el in list(doc.element.body):
    if el.tag!=qn('w:sectPr'):doc.element.body.remove(el)

def ptext(d,text,style=None,size=None,bold=False):
    p=d.add_paragraph(text,style)
    if size:
        for r in p.runs:r.font.size=Pt(size)
    if bold:
        for r in p.runs:r.bold=True
    return p

def update_cell(c,text,bold=False,left=False):
    c.text=text
    for p in c.paragraphs:
        p.alignment=WD_ALIGN_PARAGRAPH.LEFT if left else WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(3)
        p.paragraph_format.line_spacing=1.05
        for r in p.runs:r.font.name='Arial';r.font.size=Pt(10);r.bold=bold;r.font.color.rgb=RGBColor(0,0,0)

for part in ['A','B']:
    if part=='B':doc.add_page_break()
    title=ptext(doc,TABNAME+(' (continued)' if part=='B' else ''),'Title',size=12,bold=True)
    for r in title.runs:r.font.name='Arial';r.font.color.rgb=RGBColor(0,0,0)
    ptext(doc,'A. Distribution of early organ support within each trajectory class' if part=='A' else 'B. MAKE30 events within each organ support and trajectory-class group',size=10,bold=True)
    doc.element.body.insert(len(doc.element.body)-1,copy.deepcopy(original_table))
    t=doc.tables[-1]
    hdr=['Cohort','Early organ support']+[cl+('\nn (%)' if part=='A' else '\nEvents/N (%)') for cl in classes]+['Overall'+('\nn (%)' if part=='A' else '\nEvents/N (%)')]
    widths=[1.12,2.25,1.82,1.82,1.82,1.82]
    for i,c in enumerate(t.columns):c.width=Inches(widths[i])
    for j,c in enumerate(t.rows[0].cells):update_cell(c,hdr[j],True);c.width=Inches(widths[j])
    for i,db in enumerate(dbs):
        update_cell(t.cell(1+i*4,0),labels[db],True)
        for s in range(4):
            row=t.rows[1+i*4+s]
            update_cell(row.cells[1],supports[s],left=True)
            for j,cl in enumerate(classes+['Overall'],2):
                z=summary[(summary.cohort==db)&(summary.support==s)&(summary['class']==cl)].iloc[0]
                val=f'{int(z.n):,} ({z.within_class_support_percent:.1f})' if part=='A' else f'{int(z.events):,}/{int(z.n):,} ({z.event_percent:.1f})'
                update_cell(row.cells[j],val)
            for j,c in enumerate(row.cells):c.width=Inches(widths[j])
    note=('Values are n (%); percentages use the total number in the corresponding trajectory class as the denominator. Overall percentages use the cohort total. ' if part=='A' else 'Values are MAKE30 events/patients (%); percentages use the corresponding organ-support and trajectory-class group as the denominator. ')
    note+='Support is recorded within 24 hours after sepsis onset. Neither means no recorded vasopressor use or mechanical ventilation; both means both were recorded during that window, without requiring simultaneous use. MAKE30, major adverse kidney events within 30 days.'
    ptext(doc,note,size=9)
doc.save(OUT/(TABNAME+'.docx'))
assert hashlib.sha256(template.read_bytes()).hexdigest()==template_hash

legend=('Trajectory classes, early organ support and MAKE30 across three cohorts. Panels A–C show MIMIC-IV, eICU-CRD and SICdb, respectively. '
        'Ribbons connect trajectory class, organ support recorded within 24 hours after sepsis onset, and MAKE30 status; widths represent patient numbers and colours retain class identity. '
        'Each cohort is scaled separately, and node percentages use its total sample size. Vaso, vasopressors; MV, mechanical ventilation; neither, neither recorded support; both, both recorded during the window, without requiring simultaneous use. '
        'Support and trajectories are measured during the same window; connections describe patient distributions rather than treatment sequences or effects. MAKE30, major adverse kidney events within 30 days.')
leg=Document(ROOT/'JoGH_Submission_20261019/Figure legends.docx')
for el in list(leg.element.body):
    if el.tag!=qn('w:sectPr'):leg.element.body.remove(el)
leg.styles['Normal'].font.name='Arial';leg.styles['Normal'].font.size=Pt(11)
leg.styles['Title'].font.color.rgb=RGBColor(0,0,0)
ptext(leg,'Figure S5 legend','Title');ptext(leg,legend)
leg.save(OUT/'Figure S5 legend.docx')
(OUT/'Figure S5 legend.txt').write_text(legend,encoding='utf-8')

# Compact bilingual interpretation. Descriptive comparisons are not treatment effects.
both=summary[(summary.support==3)&summary['class'].isin(classes)].set_index(['cohort','class'])
lines=['早期器官支持与MAKE30分析结果','',
       '分析沿用已固定的MAKE30分析人群：MIMIC-IV 9,758例，eICU-CRD 5,709例，SICdb 3,635例。现有机械通气及升压药字段均为完整0/1记录，无额外支持字段排除。',
       '按sepsis onset后24小时内记录的升压药、机械通气构成四组：均未记录、仅升压药、仅机械通气、两者均记录。两者均记录不要求同时使用；KRT未用于分组。',
       '', '主要发现',
       '1. 三库、四种支持组合共12个分层中，观测MAKE30发生率均为C2 > C3 > C1。这个结果描述粗发生率，不是每一对比较均已达到统计学显著。',
       '2. 两类支持均有记录的患者中：']
for db in dbs:
    vals=[both.loc[(db,cl)] for cl in classes]
    lines.append(f"   {labels[db]}："+'；'.join(f'{cl} {int(z.events):,}/{int(z.n):,} ({z.event_percent:.1f}%)' for cl,z in zip(classes,vals))+'。')
lines+=['3. 支持组合与类别的关系并非三库完全相同。MIMIC-IV中，两类支持均有记录的比例为C1 50.4%、C2 44.7%、C3 26.2%，C1仍保持最低MAKE30发生率。eICU-CRD相应为23.5%、33.8%、18.1%；SICdb为32.1%、40.9%、37.9%。',
        '4. C1与C3的结局差异在四种相同支持组合内均可观察到。这使两者的生理表型差异具有更具体的临床背景，也说明两项二元支持记录不能完整概括类别间的结局差异。',
        '5. SICdb仅机械通气组较小：C1 45例、C2 23例、C3 96例。各组人数及Wilson 95%区间保留于汇总CSV，解释这些组时应关注估计精度。',
        '', '英文结果段落',
        'Early organ support patterns varied across trajectory classes and cohorts (Figure S5 and Table S6). Within each of the four support categories, observed MAKE30 incidence consistently followed the ordering C2 > C3 > C1 in all three cohorts. Among patients with both vasopressor use and mechanical ventilation recorded, MAKE30 incidence in C1, C2 and C3 was 9.7%, 49.8% and 34.1% in MIMIC-IV; 25.7%, 55.6% and 42.5% in eICU-CRD; and 18.1%, 45.8% and 25.1% in SICdb. In MIMIC-IV, both supports were more frequent in C1 than in C2 or C3, despite C1 having the lowest MAKE30 incidence. Thus, the relationship between class membership and outcomes was not captured by the presence of these two support modalities alone.',
        '', '英文方法段落',
        'We described the joint distribution of trajectory class, organ support recorded during the first 24 hours after sepsis onset, and MAKE30. Support was classified as neither vasopressors nor mechanical ventilation, vasopressors only, mechanical ventilation only, or both. The both category required records of both modalities within the window, without requiring temporal overlap. We summarised support use as within-class proportions and MAKE30 as events/patients within each class–support group. This descriptive analysis used observed binary support indicators and the frozen MAKE30 analysis cohorts.',
        '', '解释范围',
        '本分析为描述性比较。未新增调整回归、交互检验或治疗效果估计；不改变原Model 1–3和原九项交互作用检验。用于解释表型、早期支持背景和结局之间的联系。',
        '', '文件及重现',
        'Figure S5为当前JoGH投稿编号，相当于Shen原文Figure 6的展示结构；配套Table S6有两部分（支持分布及MAKE30发生率）。',
        '先运行analyse_and_plot.R完成分析和SVG绘制，再运行build_outputs.py生成PNG、600 dpi TIFF、Word表和图注。所有CSV输出均为聚合统计，不输出新增患者级文件。']
(OUT/'结果说明_中英文.txt').write_text('\n'.join(lines),encoding='utf-8')

contract=f'''# Table layout contract
Reference: {template}
SHA256: {template_hash}
Reference render: manuscript_20261019/submission_build/qa_revised/Table S2. MAKE30 and component outcomes by trajectory class_1.png
Page: A4 landscape, top/bottom 0.4722 in, left/right 0.5118 in. Preserve the reference styles, section and footer.
Title: source Title style. Six columns, 13 rows per table; first column merges each four-row cohort block; second block shaded F5F5F5.
Slots: replace title, subtitle, all cell text and footnote. Reuse the table twice, separated by a page break. Keep the source three-line border system.
Typography: Arial 10 pt cells, 9 pt notes. Center numerical cells and cohort names; left align support names. Table widths 1.12, 2.25, 1.82, 1.82, 1.82, 1.82 inches.
Two pages: Part A support n within class; Part B MAKE30 events within class/support. No individual-level data.
Permitted deviation: longer support names, added second page; no study estimates changed.
'''
(QA/'artifact.md').write_text(contract,encoding='utf-8')
manifest={'figure':FIGNAME,'table':TABNAME,'cohort_n':dict(zip(audit.cohort,audit.analysis_n.astype(int))),
          'event_counts':dict(zip(audit.cohort,audit.make30_events.astype(int))),
          'all_12_support_strata_have_C2_above_C3_above_C1':True,'new_regression_models':False,
          'source_table_unchanged':True,'class_colours':{'C1':'#0D4483','C2':'#DD1215','C3':'#39A08F'},
          'figure_legend':legend}
(OUT/'Analysis_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
print('Verified all 72 path counts independently; built figure exports and Word outputs.')
