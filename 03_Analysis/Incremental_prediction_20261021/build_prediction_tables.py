"""Create native Word tables from the new, paired prediction comparisons."""
from pathlib import Path
from copy import deepcopy
import csv
from docx import Document
from docx.shared import Inches, Pt, Cm
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "word_staging"
OUT.mkdir(exist_ok=True)
COHORTS = {"MIMIC": "MIMIC-IV", "eICU": "eICU-CRD", "SICdb": "SICdb"}

def load(name):
    with (ROOT / name).open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))

def f3(x):
    return f"{float(x):.3f}"

def ci(r, value, low, high):
    return f"{f3(r[value])}\n({f3(r[low])}–{f3(r[high])})"

def pval(x):
    x = float(x)
    return "<0.001" if x < .001 else f"{x:.3f}"

def add_border(cell, side, color="000000", size="6"):
    pr = cell._tc.get_or_add_tcPr()
    borders = pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        pr.append(borders)
    edge = OxmlElement("w:" + side)
    edge.set(qn("w:val"), "single")
    edge.set(qn("w:sz"), size)
    edge.set(qn("w:color"), color)
    borders.append(edge)

def make_doc(title, subtitle, headers, widths, values, notes, filename):
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21)
    sec.top_margin = sec.bottom_margin = Cm(1.5)
    sec.left_margin = sec.right_margin = Cm(1.4)
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(9)
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.line_spacing = 1.08
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run(title); r.bold = True; r.font.size = Pt(12)
    p = doc.add_paragraph(subtitle)
    p.paragraph_format.space_after = Pt(12)
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for col, width in zip(table.columns, widths):
        col.width = Cm(width)
    for cell, label, width in zip(table.rows[0].cells, headers, widths):
        cell.width = Cm(width)
        cell.text = label
        add_border(cell, "top", size="9")
        add_border(cell, "bottom", size="5")
        for r in cell.paragraphs[0].runs:
            r.bold = True
    trpr = table.rows[0]._tr.get_or_add_trPr()
    trpr.append(OxmlElement("w:tblHeader"))
    for i, values_row in enumerate(values):
        row = table.add_row()
        for j, (cell, value, width) in enumerate(zip(row.cells, values_row, widths)):
            cell.width = Cm(width)
            cell.text = value
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if i in [2, 3]:
                shd = OxmlElement("w:shd")
                shd.set(qn("w:fill"), "F3F3F3")
                cell._tc.get_or_add_tcPr().append(shd)
            if i == len(values) - 1:
                add_border(cell, "bottom", size="9")
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    for row in table.rows:
        for j, cell in enumerate(row.cells):
            pr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for side, value in [("top", "130"), ("bottom", "130"), ("left", "60"), ("right", "60")]:
                m = OxmlElement("w:" + side)
                m.set(qn("w:w"), value); m.set(qn("w:type"), "dxa")
                margins.append(m)
            pr.append(margins)
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.keep_with_next = False
                for run in p.runs:
                    run.font.name = "Arial"; run.font.size = Pt(9)
    for start in [1, 3, 5]:
        merged = table.cell(start, 0).merge(table.cell(start+1, 0))
        merged.text = values[start-1][0]
        merged.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in merged.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs: run.bold = True
    for i, note in enumerate(notes):
        p = doc.add_paragraph(note)
        p.paragraph_format.space_before = Pt(8 if i == 0 else 4)
        for run in p.runs: run.font.size = Pt(8.5)
    doc.save(OUT / filename)

a = load("Apparent_incremental_performance.csv")
values = []
for r in a:
    base = "SOFA" if r["comparison"] == "SOFA" else "SOFA + " + r["additional_score"]
    values.append([
        COHORTS[r["dataset"]], base, f'{int(r["events"]):,}/{int(r["n"]):,}',
        f3(r["auc_base"]), f3(r["auc_augmented"]),
        ci(r, "delta_auc", "delta_auc_low", "delta_auc_high"),
        pval(r["p_delong"]), pval(r["q_BH_six_AUC_comparisons"]),
        ci(r, "continuous_nri", "nri_low", "nri_high"),
        ci(r, "idi", "idi_low", "idi_high"),
    ])
make_doc(
    "Table S6. Incremental discrimination and reclassification for MAKE30",
    "Comparison of severity-score models before and after adding trajectory class",
    ["Cohort", "Base model", "Events/N", "Base\nAUC", "With class\nAUC", "ΔAUC\n(95% CI)", "DeLong\nP", "FDR q", "Continuous NRI\n(95% CI)", "IDI\n(95% CI)"],
    [2.5, 2.6, 2.1, 1.65, 1.8, 3.35, 1.8, 1.8, 4.1, 4.1],
    values,
    [
        "Logistic risk models were fitted separately within each cohort using continuous scores and C1–C3 membership. Paired comparisons used the same patients with complete SOFA and APS (MIMIC-IV/eICU-CRD) or SAPS 3 (SICdb) scores. AUCs are apparent estimates; cross-validation is reported in Table S7. eICU-CRD excluded 500 patients without APS.",
        "ΔAUC is the change after adding class, with paired DeLong confidence intervals and P values. FDR q values use Benjamini–Hochberg correction across the six AUC comparisons. Continuous NRI and IDI intervals use 1,000 paired patient bootstrap samples with model refitting. AUC, area under the receiver operating characteristic curve; NRI, net reclassification improvement; IDI, integrated discrimination improvement; CI, confidence interval.",
    ],
    "Table S6. Incremental discrimination and reclassification for MAKE30.docx",
)

b = load("OOF_paired_comparisons_table_ready.csv")
values = []
for r in b:
    base = "SOFA" if r["comparison"] == "SOFA" else "SOFA + " + ("SAPS 3" if r["dataset"] == "SICdb" else "APS")
    values.append([
        COHORTS[r["dataset"]], base,
        f3(r["auc_base_cv"]), f3(r["auc_augmented_cv"]), f3(r["delta_auc_cv"]),
        f3(r["brier_base"]), f3(r["brier_augmented"]),
        f3(r["calibration_slope_base"]), f3(r["calibration_slope_augmented"]),
    ])
make_doc(
    "Table S7. Internally validated performance of MAKE30 risk models",
    "Repeated stratified cross-validation within each cohort",
    ["Cohort", "Base model", "Base\nAUC", "With class\nAUC", "ΔAUC", "Base\nBrier score", "With class\nBrier score", "Base\nCalibration slope", "With class\nCalibration slope"],
    [2.5, 3.1, 2.25, 2.5, 2.0, 2.8, 2.8, 3.8, 4.05],
    values,
    [
        "Samples and model definitions match Table S6. Ten-fold stratified cross-validation was repeated ten times using identical folds for all models. AUCs are event–nonevent-pair-weighted averages of within-fold AUCs, averaged across repeats. Brier scores and calibration slopes use each patient's average held-out probability. Lower Brier scores indicate better accuracy; a calibration slope of 1 is ideal.",
        "Calibration and decision curves are shown in Figures S2 and S3. These analyses internally validate the severity-score risk models fitted within each cohort; trajectory classes retain the fixed development-model assignments. AUC, area under the receiver operating characteristic curve.",
    ],
    "Table S7. Internally validated performance of MAKE30 risk models.docx",
)
print("Created two native Word prediction tables.")
