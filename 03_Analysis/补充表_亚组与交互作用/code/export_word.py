from pathlib import Path
import pandas as pd
from docx import Document
from docx.shared import Pt,Mm,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
P=Path(__file__).resolve().parent.parent
r=pd.read_csv(P/'pooled_subgroup_RR.csv');q=pd.read_csv(P/'pooled_interaction_FDR.csv');c=pd.read_csv(P/'observed_events_and_N.csv')
note=(P/'table_legend.txt').read_text(encoding='utf-8').strip()
out=P.parent/'table2';out.mkdir(exist_ok=True)
labels={'Age':['<65 years','≥65 years'],'Sex':['Female','Male'],'Diabetes':['No','Yes']}
widths=[68,82,76,76,76,115,115,79,67]
for db in ['MIMIC','eICU','SICdb']:
 doc=Document()
 for style in doc.styles:
  for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
 s=doc.sections[0];s.page_width=Mm(297);s.page_height=Mm(210);s.left_margin=s.right_margin=Mm(15);s.top_margin=s.bottom_margin=Mm(15)
 for name in ['Normal','Title']:
  st=doc.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0);st.font.size=Pt(10 if name=='Normal' else 12);st.paragraph_format.space_after=Pt(0)
 title=('Table 2' if db=='MIMIC' else 'Supplementary Table Sx')+'  Subgroup associations of trajectory classes with MAKE30'
 t=doc.add_paragraph(title,'Title');t.runs[0].bold=True;t.paragraph_format.space_after=Pt(5)
 t=doc.add_paragraph({'MIMIC':'MIMIC-IV','eICU':'eICU-CRD','SICdb':'SICdb'}[db]+' | Model 3 adjusted analyses');t.runs[0].bold=True;t.paragraph_format.space_after=Pt(10)
 tb=doc.add_table(rows=7,cols=9);tb.autofit=False;tb.alignment=WD_TABLE_ALIGNMENT.CENTER
 # scale to available 267 mm
 widths=[x*756/754 for x in [68,82,76,76,76,115,115,79,67]]
 for col,w in zip(tb.columns,widths):col.width=Pt(w)
 headers=['Subgroup','Stratum','C1\nEvents/N','C2\nEvents/N','C3\nEvents/N','C2 vs C1\nRR (95% CI)','C3 vs C1\nRR (95% CI)','P for\ninteraction','FDR q']
 for cell,h in zip(tb.rows[0].cells,headers):cell.text=h
 for idx,(sg,levels) in enumerate(labels.items()):
  first=1+2*idx;a=q[(q.dataset==db)&(q.subgroup==sg)].iloc[0]
  for g in [0,1]:
   vals=[sg if g==0 else '',levels[g]]
   for cl in ['C1','C2','C3']:
    z=c[(c.dataset==db)&(c.subgroup==sg)&(c.stratum.astype(str)==str(g))&(c['class']==cl)].iloc[0];vals.append(f'{z.events:,.0f}/{z.n:,.0f}')
   for cl in ['C2','C3']:
    z=r[(r.dataset==db)&(r.subgroup==sg)&(r.stratum==g)&(r.comparison==cl+' vs C1')].iloc[0];vals.append(f'{z.RR:.2f} ({z.low:.2f}–{z.high:.2f})')
   vals+=([f'{a.p_interaction:.3f}',f'{a.q_BH:.3f}'] if g==0 else ['',''])
   for cell,v in zip(tb.rows[first+g].cells,vals):cell.text=v
  for j in [0,7,8]:
   v=tb.cell(first,j).text;cell=tb.cell(first,j).merge(tb.cell(first+1,j));cell.text=v
 for i,row in enumerate(tb.rows):
  row.height=Pt(35 if i==0 else 31)
  trPr=row._tr.get_or_add_trPr();trPr.append(OxmlElement('w:cantSplit'))
  if i==0:trPr.append(OxmlElement('w:tblHeader'))
  for j,(cell,w) in enumerate(zip(row.cells,widths)):
   cell.width=Pt(w);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   for para in cell.paragraphs:
    para.alignment=WD_ALIGN_PARAGRAPH.CENTER;para.paragraph_format.space_before=Pt(0);para.paragraph_format.space_after=Pt(0)
    for run in para.runs:run.font.size=Pt(10);run.bold=(i==0 or j==0)
   borders=OxmlElement('w:tcBorders')
   for side in ['top','bottom','left','right']:
    el=OxmlElement('w:'+side);el.set(qn('w:val'),'single' if (i==0 and side in ['top','bottom']) or (i==6 and side=='bottom') else 'nil');el.set(qn('w:sz'),'6');el.set(qn('w:color'),'000000');borders.append(el)
   cell._tc.get_or_add_tcPr().append(borders)
 # bottom border on merged final-group cells needs explicit application
 for j in [0,7,8]:
  cell=tb.cell(5,j);b=cell._tc.get_or_add_tcPr().find(qn('w:tcBorders'));e=b.find(qn('w:bottom'));e.set(qn('w:val'),'single')
 para=doc.add_paragraph(note);para.paragraph_format.space_before=Pt(7)
 for run in para.runs:run.font.size=Pt(9)
 target=out/'Table2.docx' if db=='MIMIC' else P/f'Supplementary_Table_{db}.docx'
 # Set table-wide bottom border including vertically merged cells.
 for row in tb.rows:
  for tc in row._tr.findall(qn('w:tc')):
   pr=tc.get_or_add_tcPr()
   for b in pr.findall(qn('w:tcBorders')):
    for el in list(b):
     if el.tag==qn('w:bottom'):b.remove(el)
 borders=OxmlElement('w:tblBorders')
 el=OxmlElement('w:bottom');el.set(qn('w:val'),'single');el.set(qn('w:sz'),'6');el.set(qn('w:color'),'000000');borders.append(el);tb._tbl.tblPr.append(borders)
 # Header separator remains explicit.
 for cell in tb.rows[0].cells:
  borders=cell._tc.get_or_add_tcPr().find(qn('w:tcBorders'));el=OxmlElement('w:bottom');el.set(qn('w:val'),'single');el.set(qn('w:sz'),'6');el.set(qn('w:color'),'000000');borders.append(el)
 doc.save(target);print(target)
