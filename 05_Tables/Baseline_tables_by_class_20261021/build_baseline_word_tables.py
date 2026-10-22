"""Update Table 1 / S1 and clone class-stratified baseline Word tables.

Only document.xml is replaced in each retained DOCX package. Other parts,
including all styles, relationships and section furniture, are preserved.
"""
from copy import deepcopy
from hashlib import sha256
from pathlib import Path
from zipfile import ZipFile
import json
import pandas as pd
from docx import Document
from docx.oxml.ns import qn
from docx.table import _Cell
from lxml import etree

WORK = Path("Z:/Research/MSI-1/analysis_updates/Baseline_tables_by_class_20261021")
STAGE = WORK / "word_staging"
STAGE.mkdir(exist_ok=True)
QA = Path(Path("Z:/Research/MSI-1/analysis_updates/active_extension_qa.txt").read_text(
    encoding="utf-8-sig").strip())
SUBMISSION = Path("Z:/Research/MSI-1/JoGH_Submission_20261019")
NAME1 = "Table 1. Baseline characteristics of the three study cohorts.docx"
NAMES1 = "Table S1. Between-cohort standardised differences and missingness of baseline characteristics.docx"
REF1 = QA / "formal_backup" / NAME1
REFS1 = QA / "supp_backup" / NAMES1
TARGET1 = SUBMISSION / "Formal figures and tables" / NAME1
TARGETS1 = SUBMISSION / "Supplementary materials" / NAMES1
DATA = json.loads((WORK / "baseline_tables_by_class.json").read_text(encoding="utf-8"))
MORT = {v["database"]: v for v in json.loads(
    (WORK / "new_mortality_row_for_Table1_and_S1.json").read_text(encoding="utf-8"))}
MORTLABEL = "30-day in-hospital mortality after ICU admission"
MORTNOTE = (
    "30-day in-hospital mortality is reported as events/evaluable patients (%) "
    "and was assessed from ICU admission using separately adjudicated hospital status and timing."
)
DB_NAMES = {"MIMIC": "MIMIC-IV", "eICU": "eICU-CRD", "SICdb": "SICdb"}
REPORT = []
LUNG_PATH = Path("Z:/Research/MSI-1/analysis_updates/Manuscript_results_discussion_20261021/Patient_profiles_observed_summary.csv")
LUNG_SOURCE = pd.read_csv(LUNG_PATH)


def text_paragraph(p, text):
    rpr = deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
    for node in list(p._p):
        if node.tag != qn("w:pPr"):
            p._p.remove(node)
    run = p.add_run(text)
    if rpr is not None:
        run._r.insert(0, rpr)


def text_cell(tc, table, text):
    cell = _Cell(tc, table)
    text_paragraph(cell.paragraphs[0], text)
    for p in cell.paragraphs[1:]:
        p._p.getparent().remove(p._p)


def physical_rows(doc):
    return [[[_Cell(tc, t).text for tc in tr.tc_lst] for tr in t._tbl.tr_lst]
            for t in doc.tables]


def repair_spans(doc):
    for t in doc.tables:
        ncols = len(t.columns)
        width = sum(int(c.get(qn("w:w"))) for c in t._tbl.tblGrid)
        for tr in t._tbl.tr_lst:
            if len(tr.tc_lst) == 1:
                tcpr = tr.tc_lst[0].get_or_add_tcPr()
                span = tcpr.find(qn("w:gridSpan"))
                if span is not None:
                    span.set(qn("w:val"), str(ncols))
                tcw = tcpr.find(qn("w:tcW"))
                if tcw is not None:
                    tcw.set(qn("w:w"), str(width))


def update_outcomes(doc):
    for t in doc.tables:
        for tr in t._tbl.tr_lst:
            if len(tr.tc_lst) == 1 and "Outcomes (frozen MAKE30 definitions)" in _Cell(tr.tc_lst[0], t).text:
                text_cell(tr.tc_lst[0], t, "Outcomes")


def append_row(table, texts):
    cloned = deepcopy(table._tbl.tr_lst[-1])
    assert len(cloned.tc_lst) == len(texts)
    table._tbl.append(cloned)
    for tc, txt in zip(cloned.tc_lst, texts):
        text_cell(tc, table, txt)


def save_package(reference, doc, target):
    xml = etree.tostring(doc._element, xml_declaration=True, encoding="UTF-8", standalone=True)
    with ZipFile(reference) as src, ZipFile(target, "w") as dest:
        for item in src.infolist():
            dest.writestr(item, xml if item.filename == "word/document.xml" else src.read(item.filename))
    with ZipFile(reference) as src, ZipFile(target) as dest:
        preserved = [n for n in src.namelist() if n != "word/document.xml"]
        assert all(src.read(n) == dest.read(n) for n in preserved)
    REPORT.append({"file": str(target), "sha256": sha256(target.read_bytes()).hexdigest(),
                   "preserved_package_parts": len(preserved),
                   "rows": [len(t.rows) for t in doc.tables],
                   "sections": len(doc.sections),
                   "portrait": all(s.page_height > s.page_width for s in doc.sections)})


def verify_existing_cells(reference, newdoc):
    old = physical_rows(Document(reference))
    now = physical_rows(newdoc)
    assert len(old) == len(now) == 2
    for i, rows in enumerate(old):
        for j, cells in enumerate(rows):
            expected = [x.replace("Outcomes (frozen MAKE30 definitions)", "Outcomes") for x in cells]
            assert now[i][j] == expected, (i, j, expected, now[i][j])


def write_contract():
    references = []
    for p in (REF1, REFS1):
        with ZipFile(p) as z:
            inventory = [{"path": n, "bytes": len(z.read(n)),
                          "sha256": sha256(z.read(n)).hexdigest(),
                          "status": "editable" if n == "word/document.xml" else "preserve-only"}
                         for n in z.namelist()]
        references.append({"path": str(p), "sha256": sha256(p.read_bytes()).hexdigest(),
                           "package": inventory})
    (QA / "baseline_tables_qa" / "reference_package_inventory.json").write_text(
        json.dumps(references, ensure_ascii=False, indent=2), encoding="utf-8")
    contract = f"""# Baseline table template contract

Reference Table 1: {REF1}
Reference S1: {REFS1}
Both references have been rendered and all four reference pages visually inspected.
Reference hashes and package inventories: reference_package_inventory.json.
Reference pages: reference_table1/pages and reference_s1/pages.

Page system: one A4 portrait section (11909 x 16834 twips); top/bottom 864 twips,
left/right 1080 twips. No visible headers or footers. Two manually separated pages.

Typography: title and continued title Arial 10 pt bold black; table headers and
section titles Arial 9 pt bold; cells Arial 9 pt black; retained single line spacing,
2 pt before/after cell paragraphs, 22-twip top/bottom cell padding, 35-twip right padding.
Footnote inherits retained paragraph style and is visually smaller than table text.

Tables: two native Word tables, fixed layout; Table1 grid 3456/2088/2088/2088 twips;
S1 grid 3024/1080/1080/1512/1512/1512 twips. Header top/bottom black 0.5 pt rules.
Header repeats, regular rows do not split; section labels are merged and bold.
First table contains demographics, vital signs, scores and laboratory values.
Second table contains comorbidities, treatment, MSI and outcomes.

Editable slots: Table1 outcomes section label; final added mortality row; note sentence.
S1 same section label; final added missingness/SMD row; supporting note.
New class tables clone Table1; replace titles, database column headers with C1/C2/C3
and fixed class sizes, each numeric cell with the corresponding aggregate JSON value,
and use a short table-specific note.
Retain all old Table1 and S1 numeric/data text exactly. No images/fields/content controls.

Intentional structural repair: merged section-label rows still have gridSpan=9 from
an earlier nine-column layout; normalize to actual 4 or 6 columns without changing text.
No other package member, relationship, style, numbering or section geometry changes.
All final pages must be rendered through canonical render_docx.py using the Word-export
and PDFium adapter because this Windows dependency bundle has no LibreOffice.
"""
    (QA / "baseline_tables_qa" / "artifact.md").write_text(contract, encoding="utf-8")


def main():
    write_contract()
    doc = Document(REF1)
    repair_spans(doc)
    update_outcomes(doc)
    append_row(doc.tables[-1], [MORTLABEL] + [MORT[d]["display"] for d in ["MIMIC", "eICU", "SICdb"]])
    note = doc.paragraphs[-1]
    text_paragraph(note, note.text + " " + MORTNOTE)
    verify_existing_cells(REF1, doc)
    save_package(REF1, doc, TARGET1)
    for i, table in enumerate(doc.tables, 1):
        (STAGE / f"Table1_table_{i}.xml").write_text(table._tbl.xml, encoding="utf-8")
    (STAGE / "Table1_note.xml").write_text(note._p.xml, encoding="utf-8")
    (STAGE / "Table1_note.txt").write_text(note.text, encoding="utf-8")
    (STAGE / "Table1_document.xml").write_bytes(
        etree.tostring(doc._element, xml_declaration=True, encoding="UTF-8", standalone=True))

    s1 = Document(REFS1)
    repair_spans(s1)
    update_outcomes(s1)
    append_row(s1.tables[-1], [MORTLABEL, MORT["eICU"]["binary_smd_display"],
                              MORT["SICdb"]["binary_smd_display"]]
               + [MORT[d]["unknown_display"] for d in ["MIMIC", "eICU", "SICdb"]])
    s1note = s1.paragraphs[-1]
    text_paragraph(s1note, s1note.text + " For 30-day in-hospital mortality, missingness is shown as n/cohort N (%), and SMDs use evaluable outcomes.")
    verify_existing_cells(REFS1, s1)
    save_package(REFS1, s1, TARGETS1)

    for table_number, db in enumerate(["MIMIC", "eICU", "SICdb"], 2):
        d = Document(TARGET1)
        cohort = DATA["cohorts"][db]
        label_to_data = {r["label"]: r for r in cohort["rows"]}
        title = f"Table S{table_number}. Baseline characteristics by trajectory class in {DB_NAMES[db]}"
        for p in d.paragraphs:
            if p.text.startswith("Table 1."):
                text_paragraph(p, title + (" (continued)" if "(continued)" in p.text else ""))
        for t in d.tables:
            for idx, tr in enumerate(t._tbl.tr_lst):
                if len(tr.tc_lst) == 1:
                    continue
                if idx == 0:
                    for k, c in enumerate(["C1", "C2", "C3"], 1):
                        text_cell(tr.tc_lst[k], t, f"{c}\n(n = {cohort['class_n'][c]})")
                else:
                    label = _Cell(tr.tc_lst[0], t).text
                    record = label_to_data[label]
                    for k, c in enumerate(["C1", "C2", "C3"], 1):
                        text_cell(tr.tc_lst[k], t, record["values"][c]["display"])
        lung_variable = "precondition_lung_disease" if db == "SICdb" else "chronic_pulmonary_disease"
        lung_label = "Recorded prior lung disease" if db == "SICdb" else "Chronic pulmonary disease"
        lung_rows = LUNG_SOURCE[
            LUNG_SOURCE.cohort.eq("Full trajectory cohort")
            & LUNG_SOURCE.dataset.eq(db)
            & LUNG_SOURCE.variable.eq(lung_variable)].set_index("class")
        assert len(lung_rows) == 3
        lung_values = {}
        for c in ["C1", "C2", "C3"]:
            z = lung_rows.loc[c]
            assert int(z.class_n) == cohort["class_n"][c]
            assert int(z.observed_n) == cohort["class_n"][c]
            lung_values[c] = {"display": f"{int(z.positive_n)} ({z.percent_observed:.1f})"}
        label_to_data[lung_label] = {"label": lung_label, "values": lung_values}
        for t in d.tables:
            for tr in list(t._tbl.tr_lst):
                if len(tr.tc_lst) == 4 and _Cell(tr.tc_lst[0], t).text == "Congestive heart failure":
                    cloned = deepcopy(tr)
                    tr.addnext(cloned)
                    for tc, txt in zip(cloned.tc_lst, [lung_label] + [lung_values[c]["display"] for c in ["C1", "C2", "C3"]]):
                        text_cell(tc, t, txt)
        short_note = (
            "Data are median [Q1, Q3] or n (%), using available observations. "
            "Variable definitions and handling follow Table 1. Dynamic measurements "
            "are averaged over the first 24 hours after sepsis onset; the MSI summaries "
            "are defined in the row labels. BMI is calculated from weight and height "
            "when a direct value is unavailable. N.A. indicates an unavailable measure. "
            + MORTNOTE + " APS, Acute Physiology Score; MAKE30, major adverse kidney events within 30 days; "
            "MSI, modified shock index; KRT, kidney replacement therapy."
        )
        text_paragraph(d.paragraphs[-1], short_note)
        target = STAGE / f"{title}.docx"
        save_package(REF1, d, target)
        # Verify every numerical class cell against the aggregate data.
        for t in d.tables:
            for idx, tr in enumerate(t._tbl.tr_lst):
                if idx == 0 or len(tr.tc_lst) == 1:
                    continue
                label = _Cell(tr.tc_lst[0], t).text
                assert [_Cell(tc, t).text for tc in tr.tc_lst[1:]] == [
                    label_to_data[label]["values"][c]["display"] for c in ["C1", "C2", "C3"]]
    manifest = []
    for item in REPORT:
        p = Path(item["file"])
        key = p.stem.split(".")[0].replace(" ", "_")
        manifest.append({"docx": str(p), "pdf": str(QA / "baseline_tables_qa" / key / f"{key}.pdf"),
                         "render_dir": str(QA / "baseline_tables_qa" / key / "pages")})
    (WORK / "final_render_manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    (WORK / "word_validation.json").write_text(
        json.dumps({"files": REPORT, "old_Table1_and_S1_cells_unchanged": True,
                    "all_new_class_cells_match_aggregate_JSON": True,
                    "additional_lung_rows_match_existing_full_trajectory_profile_CSV": str(LUNG_PATH),
                    "existing_other_package_parts_unchanged": True},
                   ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(REPORT, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
