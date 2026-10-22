from pathlib import Path
import json,copy
from reportlab.platypus import SimpleDocTemplate,Table,TableStyle,Paragraph,Spacer,PageBreak
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.pagesizes import A4,landscape
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import pypdfium2 as pdfium
ROOT=Path(__file__).resolve().parent.parent
OUT=ROOT/'derived'
rows=json.loads((OUT/'Table1_rows.json').read_text(encoding='utf-8'))
headers=['Variable','MIMIC-IV\n(n = 9804)','eICU-CRD\n(n = 5839)','SICdb\n(n = 3831)','SMD¹','SMD²','Missing\n(MIMIC)','Missing\n(eICU)','Missing\n(SICdb)']
keys=['label','MIMIC','eICU','SICdb','smd_eICU','smd_SICdb','missing_MIMIC','missing_eICU','missing_SICdb']
widths=[188,100,100,100,43,43,67,67,67]
notes=[
'Data are median [Q1, Q3] or n (%). Missing values are n (% of the full cohort). For comorbidities, missingness is reported BEFORE coding unrecorded values as 0; percentages then use the full cohort.',
'SMD¹: eICU versus MIMIC; SMD²: SICdb versus MIMIC. Absolute SMD uses means and pooled SDs for continuous variables, and proportions for binary variables.',
'For repeatedly measured vital signs and current laboratory variables, the mean within 24 hours after sepsis onset is used. SICdb laboratory source priority: central laboratory, then explicitly arterial blood gas for sodium, potassium, chloride, bicarbonate and glucose. Blood pressure: invasive, then non-invasive; MIMIC temperature: Celsius, then converted Fahrenheit. Means are calculated within the selected source; channel means are not averaged together.',
'When BMI is not directly available, it is calculated from weight and height as weight (kg) divided by height (m) squared.',
'N.A. indicates that a variable is unavailable for that database in the current analysis, rather than patient-level missingness. SMD is not calculated when either database lacks the variable; unavailable SMDs are shown as —. APS, Acute Physiology Score.',
'Charlson Comorbidity Index includes age points and is derived from recorded ICD diagnoses; unavailable diagnostic codes yield a missing score.',
'SOFA is the frozen sepsis-onset score; other severity scores are admission/database-specific. Baseline creatinine and eGFR, treatment windows, MSI and MAKE30 outcomes retain their frozen definitions. MAKE30 outcomes are unavailable when the frozen evaluation window is invalid. Ischemic heart disease codes include acute myocardial infarction.'
]
pdfmetrics.registerFont(TTFont('Light','C:/Windows/Fonts/calibril.ttf'));pdfmetrics.registerFont(TTFont('Bold','C:/Windows/Fonts/calibrib.ttf'))
sty=ParagraphStyle('cell',fontName='Light',fontSize=8,leading=9)
bsty=ParagraphStyle('bold',parent=sty,fontName='Bold')
nsty=ParagraphStyle('note',parent=sty,fontSize=7,leading=8.5)
split=next(i for i,r in enumerate(rows) if r['section']=='Comorbidities');groups=[rows[:split],rows[split:]]
def content(group):
 result=[('header',headers)];prev=None
 for r in group:
  if r['section']!=prev:result.append(('section',[r['section']]+['']*8));prev=r['section']
  result.append(('data',[r[k] for k in keys]))
 return result
story=[]
for page,group in enumerate(groups):
 if page:story.append(PageBreak())
 story.append(Paragraph('Table 1'+(' (continued)' if page else '')+'  Baseline characteristics of the three study cohorts',bsty));story.append(Spacer(1,6))
 expanded=content(group)
 cells=[[Paragraph(str(v).replace('&','&amp;').replace('\n','<br/>'),bsty if kind!='data' else sty) for v in vals] for kind,vals in expanded]
 commands=[('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),0),('RIGHTPADDING',(0,0),(-1,-1),3),('TOPPADDING',(0,0),(-1,-1),1.1),('BOTTOMPADDING',(0,0),(-1,-1),1.1),('LINEABOVE',(0,0),(-1,0),.5,'black'),('LINEBELOW',(0,0),(-1,0),.5,'black'),('LINEBELOW',(0,-1),(-1,-1),.5,'black')]
 for i,(kind,_) in enumerate(expanded):
  if kind=='section':commands.append(('SPAN',(0,i),(-1,i)))
 t=Table(cells,colWidths=widths,repeatRows=1);t.setStyle(TableStyle(commands));story.append(t)
 if page==1:
  for note in notes:story.extend([Spacer(1,3),Paragraph(note,nsty)])
pdf=ROOT/'Table1.pdf'
SimpleDocTemplate(str(pdf),pagesize=landscape(A4),leftMargin=33,rightMargin=33,topMargin=22,bottomMargin=20).build(story)
# Clone the retained template's native table and rows, preserving its formatting.
doc=Document(ROOT/'templates/style_template.docx')
template=copy.deepcopy(doc.tables[0]._tbl)
body=doc._element.body
for child in list(body):
 if child.tag!=qn('w:sectPr'):body.remove(child)
from docx.table import Table as WordTable
for page,group in enumerate(groups):
 if page:doc.add_page_break()
 p=doc.add_paragraph('Table 1'+(' (continued)' if page else '')+'  Baseline characteristics of the three study cohorts');p.runs[0].bold=True
 el=copy.deepcopy(template);body.insert(len(body)-1,el);t=WordTable(el,doc)
 header=copy.deepcopy(t.rows[0]._tr);section=copy.deepcopy(t.rows[1]._tr);data=copy.deepcopy(t.rows[2]._tr)
 for tr in list(el.findall(qn('w:tr'))):el.remove(tr)
 for kind,vals in content(group):
  el.append(copy.deepcopy(header if kind=='header' else section if kind=='section' else data))
  rr=t.rows[-1]
  for cell,val in zip(rr.cells[:1] if kind=='section' else rr.cells,vals[:1] if kind=='section' else vals):
   p=cell.paragraphs[0];ppr=copy.deepcopy(p._p.pPr)
   for child in list(p._p):p._p.remove(child)
   if ppr is not None:p._p.append(ppr)
   run=p.add_run(str(val));run.font.name='Calibri' if kind!='data' else 'Calibri Light';run.font.size=Pt(8);run.bold=kind!='data'
  if kind=='data':
   for cell in rr.cells:
    borders=cell._tc.get_or_add_tcPr().find(qn('w:tcBorders'))
    if borders is not None:cell._tc.get_or_add_tcPr().remove(borders)
 if page==1:
  for note in notes:
   p=doc.add_paragraph(note);p.paragraph_format.space_before=Pt(3);p.paragraph_format.line_spacing=Pt(8.5)
   for run in p.runs:run.font.size=Pt(7)
doc.save(ROOT/'Table1.docx')
p=pdfium.PdfDocument(str(pdf));assert len(p)==2,len(p)
for i in range(len(p)):p[i].render(scale=1.7).to_pil().save(OUT/f'Table1_第{i+1}页.png')
print('Created 2-page PDF, editable DOCX, and previews.')
