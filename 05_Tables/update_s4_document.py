from pathlib import Path
import pandas as pd,shutil,json,hashlib
from docx import Document
from docx.shared import Pt
root=Path('Z:/Research/MSI-1');out=root/'analysis_updates/S4_Model3_20261019';res=root/'results';p=next((res/'02_Supplementary_Figures_and_Tables').glob('Table S4*.docx'))
shutil.copy2(p,out/'Previous_Table_S4.docx')
new=pd.read_csv(out/'Model3_sensitivity_pooled.csv');old=pd.read_csv(root/'历史归档/results_整理前_20261019/results/V331_MAR_MNAR敏感性_G3/V331_分类汇总_G3.csv')
d=Document(p);t=d.tables[0];rows=[]
for bi,db in enumerate(['MIMIC','eICU','SICdb']):
 for si,scenario in enumerate(['Original','MAR','MNAR']):
  idx=1+bi*3+si
  for ci,comparison in [(3,'C2 vs C1'),(4,'C3 vs C1')]:
   z=new[(new.dataset==db)&(new.comparison==comparison)&(new.scenario.str.startswith('MNAR') if scenario=='MNAR' else new.scenario.eq(scenario))]
   val=f'{z.RR.min():.2f}–{z.RR.max():.2f}' if scenario=='MNAR' else f'{z.RR.iloc[0]:.2f} ({z.lower.iloc[0]:.2f}–{z.upper.iloc[0]:.2f})'
   cell=t.cell(idx,ci);cell.paragraphs[0].runs[0].text=val
   for r in cell.paragraphs[0].runs[1:]:r.text=''
  rows.append([c.text for c in t.rows[idx].cells])
foot='Agreement is with the original classification. RRs were estimated using modified Poisson regression adjusted for the same Model 3 covariates as the main analysis. MNAR rows summarize eight scenarios: agreement ranges are across scenario means, and RR ranges are across pooled point estimates, not confidence intervals. All corresponding 95% CIs remained above 1. MAR/MNAR, missing at random/not at random.'
for para in d.paragraphs:
 if para.text.startswith('Agreement is with'):
  para.text=foot
  for r in para.runs:r.font.size=Pt(9)
d.save(p);pd.DataFrame(rows,columns=['Cohort','Scenario','Agreement','C2_RR','C3_RR']).to_csv(out/'Final_Table_S4_display.csv',index=False)
shutil.copy2(p,out/'Updated_Table_S4.docx')
code=res/'03_Final_Code/S4_Model3_update';code.mkdir(exist_ok=True)
for f in ['run_s4_model3.R','verify_s4.R','Model3_sensitivity_pooled.csv','New_vs_previous_results.csv','Final_Table_S4_display.csv','Imputation_logged_events_review.csv']:shutil.copy2(out/f,code/f)
methods='''# S4统一Model 3后的分析记录

本次仅更新敏感性情景中的协变量插补与MAKE30关联回归；原始提取、结局、冻结轨迹模型和正式分类均未改变。

每库使用原V331保存的20套轨迹插补对象，按原MAR及8种MNAR偏移规则重新计算临时类别。逐库逐次核对分类一致率和ARI，与旧版全部一致。协变量采用正文清洗数据、同一15协变量、同一PMM/二分类logistic插补方法及预测矩阵；保留结局和该情景下的临时类别作为预测变量。对每个轨迹插补及情景，运行一条独立的协变量插补链30轮，从正文已完成数据作为初始化，获取一个条件协变量插补结果。20个完整结果在log RR尺度按Rubin及Barnard–Rubin自由度合并。未将20与30套数据做笛卡尔组合，未把重复数据当独立插补。

Original行使用正文保存的全部30套协变量插补，重新计算结果与正文Model 3的RR和CI逐项匹配至1e-8内。其余情景各20套，并在每套内使用log-link Poisson、HC1协方差以及相同缩放方法。所有模型收敛且无秩亏。540个条件插补/回归已核查；MICE记录仅为BMI插补时移除线性冗余的weight_missing1辅助预测项，未删除任何分析协变量。最大Monte Carlo SE/total SE为0.1204，数值按两位小数展示。最终值及逐次系数、插补值、链均值、临时分类、种子等保存在analysis_updates/S4_Model3_20261019。

新旧差异同时涉及调整变量和配套协变量插补；不能将变化全部归因于某一个新增协变量。分类一致率不变；三库全部敏感性情景中C2/C3相对C1的95%CI下限均>1。

G3支持图表仅作为待审阅的模型选择支持材料，未自动加入正式补充编号。其34个类别对来源于保存的初筛拟合点，既非最终精修模型之间的比较，也不是新的轨迹拟合。
'''
(out/'Methods_and_validation.md').write_text(methods,encoding='utf-8');shutil.copy2(out/'Methods_and_validation.md',code/'Methods_and_validation.md')
rp=res/'README_图表说明.md';txt=rp.read_text(encoding='utf-8')
start=txt.index('4. 补充Table S4');end=txt.index('\n5.',start)
txt=txt[:start]+'4. 补充Table S4现已采用与正文相同的Model 3协变量和修正Poisson/HC1回归。Original行复现正文30套插补结果，MAR/MNAR情景结合20套轨迹插补与相应协变量条件插补估计；MNAR范围是各情景点估计范围，不是置信区间。'+txt[end:]
txt=txt.replace('本轮只汇总、命名及压缩Table 1表注，没有改动任何表格数值或重新拟合模型。','目录整理阶段仅汇总、命名及压缩Table 1表注；随后经用户授权，S4已更新为Model 3敏感性回归，正文图表数值未改变。')
txt+='\n\n## S4最新更新\n已更新为正文Model 3调整方案；原分类下的RR及95%CI与正文一致。MAR/MNAR分类一致率与旧版一致。新旧RR对照和实施记录位于03_Final_Code/S4_Model3_update。G3水平—形态支持图表暂放analysis_updates/G3_level_shape_support_20261019供审阅，未增加正式补充编号。\n'
rp.write_text(txt,encoding='utf-8')
# Record current source code without disturbing the three-folder delivery structure.
shutil.copy2(root/'analysis_updates/build_level_shape_support.py',res/'03_Final_Code/build_level_shape_support.py')
man=res/'03_Final_Code/delivery_manifest.json';lst=json.loads(man.read_text(encoding='utf-8'))
for item in lst:
 if item['output'].startswith('02_Supplementary') and 'Table S4.' in item['output']:
  item['previous_archived_source']=item.pop('archived_source',None);item['source_analysis']=str(out);item['table_cells_unchanged']=False;item['update']='Model 3 sensitivity regression';item['sha256']=hashlib.sha256(p.read_bytes()).hexdigest()
man.write_text(json.dumps(lst,ensure_ascii=False,indent=2),encoding='utf-8')
print(p)
