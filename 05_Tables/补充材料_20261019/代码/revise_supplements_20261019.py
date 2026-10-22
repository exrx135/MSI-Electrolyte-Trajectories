from pathlib import Path
import shutil, pandas as pd
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
old=Path('Z:/Research/MSI-1/results/候选补充材料_20261018')
P=Path('Z:/Research/MSI-1/results/补充材料_20261019')
for sub in ['表','图','代码','核查','数据']:(P/sub).mkdir(parents=True,exist_ok=True)
backup=Path('Z:/Research/MSI-1/results/审稿备用_模型选择与定义');backup.mkdir(exist_ok=True)
for prefix in ['TableS1','TableS2','TableS6']:
 for f in (old/'表').glob(prefix+'*'):shutil.copy2(f,backup/f.name)
for f in (old/'图').iterdir():
 if f.is_file() and '预览' not in f.name:shutil.copy2(f,P/'图'/f.name)
for a,b,stem in [(3,1,'MAKE30_components'),(4,2,'Subgroups_eICU'),(5,3,'Subgroups_SICdb')]:
 d=Document(old/'表'/f'TableS{a}_{stem}.docx')
 for p in d.paragraphs:
  for r in p.runs:r.text=r.text.replace(f'Table S{a}',f'Table S{b}')
 d.save(P/'表'/f'TableS{b}_{stem}.docx')
# Reuse approved table layout helpers only; do not execute the old analysis builder.
source=(old/'代码/build_candidate_tables.py').read_text(encoding='utf-8').split('# S1 definitions')[0]
ns={};exec(source,ns)
x=pd.read_csv(old/'数据/S6_classification.csv');rr=pd.read_csv(old/'数据/S6_archived_RR.csv')
d=ns['new']();ns['title'](d,4,'Sensitivity to missing trajectory measurements','Classification agreement and MAKE30 associations across three cohorts')
rows=[]
for db,label in [('MIMIC','MIMIC-IV'),('eICU','eICU-CRD'),('SICdb','SICdb')]:
 for scen,text in [('Original','Original'),('MAR','MAR'),('MNAR','MNAR scenarios')]:
  a=x[(x.dataset==db)&(x.scenario.str.startswith('MNAR') if scen=='MNAR' else x.scenario.eq(scen))]
  agreement=f'{a.agreement_mean.min()*100:.1f}–{a.agreement_mean.max()*100:.1f}' if scen=='MNAR' else f'{a.agreement_mean.iloc[0]*100:.1f}'
  vals=[label,text,agreement]
  for comp in ['C2 vs C1','C3 vs C1']:
   b=rr[(rr.dataset==db)&(rr.comparison==comp)&(rr.scenario.str.startswith('MNAR') if scen=='MNAR' else rr.scenario.eq(scen))]
   assert (b.lower>1).all()
   vals.append(f'{b.RR.min():.2f}–{b.RR.max():.2f}' if scen=='MNAR' else f'{b.RR.iloc[0]:.2f} ({b.lower.iloc[0]:.2f}–{b.upper.iloc[0]:.2f})')
  rows.append(vals)
t=ns['tbl'](d,['Cohort','Scenario','Same-class\nagreement (%)','C2 vs C1\nRR (95% CI) or range','C3 vs C1\nRR (95% CI) or range'],rows,[120,140,150,180,180],font=10,height=31,merges=[(0,0,2),(0,3,5),(0,6,8)])
for i in range(4,7):
 for tc in t.rows[i]._tr.findall(qn('w:tc')):
  sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F5F5F5');tc.get_or_add_tcPr().append(sh)
foot='Agreement is with the original classification, averaged across 20 imputations per scenario. MNAR rows show the range of scenario-mean agreement and pooled RR point estimates across eight scenarios, not confidence intervals. All corresponding 95% CIs remained above 1. Only missing trajectory measurements were imputed or shifted; model parameters were fixed. RRs use the existing adjustment for age, sex, SOFA, baseline eGFR, kidney disease, diabetes, mechanical ventilation and vasopressor use, distinct from the current Model 3. MAR/MNAR, missing at random/not at random.'
ns['note'](d,foot);d.save(P/'表/TableS4_Missing_measurement_sensitivity.docx')
pd.DataFrame(rows,columns=['cohort','scenario','agreement_percent','C2_RR','C3_RR']).to_csv(P/'数据/TableS4_summary.csv',index=False)
for f in ['S6_classification.csv','S6_archived_RR.csv']:shutil.copy2(old/'数据'/f,P/'数据'/f)
legend=(old/'图注与表注.md').read_text(encoding='utf-8')
(P/'图注与表注.md').write_text(legend.split('## TableS1')[0],encoding='utf-8')
# Extract current table notes directly from their editable Word originals.
with (P/'图注与表注.md').open('a',encoding='utf-8') as out:
 for p in sorted((P/'表').glob('*.docx')):
  doc=Document(p);out.write('\n\n## '+p.stem+'\n'+ '\n'.join(q.text for q in doc.paragraphs if q.text)+'\n')
(P/'请先阅读.md').write_text('当前补充材料：2幅图、4张表。均为供审阅版本。\n\nFigure S1：eICU类内状态变化；Figure S2：SICdb类内状态变化。\nTable S1：MAKE30及组成结局（原S3）；Table S2：eICU亚组（原S4）；Table S3：SICdb亚组（原S5）；Table S4：三库缺失轨迹敏感性简表（原S6）。\n\n原定义表、模型选择表和完整敏感性表保存于相邻“审稿备用_模型选择与定义”文件夹，不包含在当前合并PDF和ZIP中。未重跑模型或数据提取。\n\nS4将8个MNAR情景概括为范围；RR范围是点估计范围，不是置信区间。完整情景数据保留。三库总体结果具有稳定性，SICdb分类一致性最高；同时区分类别归属一致性和关联方向稳定性，保留实际数值。\n',encoding='utf-8')
print(P)
print(pd.DataFrame(rows).to_string(index=False))
